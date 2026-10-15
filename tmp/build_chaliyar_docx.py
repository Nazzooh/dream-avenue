import os
import sys
import tempfile

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


FACING = r"C:\Users\Sabreena\Downloads\FACING SHEETS DISSERTATION-1 (1).pdf"
PROJECT = r"C:\Users\Sabreena\Downloads\Chaliyar_River_Pollution_Project_1.pdf"
GUIDE = r"C:\Users\Sabreena\Downloads\Guide Biodata cum Certificate of Willingness.pdf"
OUTPUT = os.path.join(os.path.dirname(os.path.dirname(__file__)), "Chaliyar_River_Pollution_Project_1.docx")


def render_pages(pdf_paths, image_dir):
    page_images = []
    for pdf_path in pdf_paths:
        document = pymupdf.open(pdf_path)
        for page_number, page in enumerate(document, start=1):
            image_path = os.path.join(image_dir, f"page_{len(page_images) + 1:02d}.png")
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            pixmap.save(image_path)
            page_images.append((image_path, page.rect.width / 72, page.rect.height / 72))
        document.close()
    return page_images


def add_page(document, image_path, width_inches, height_inches, is_last):
    section = document.sections[-1]
    section.page_width = Inches(width_inches)
    section.page_height = Inches(height_inches)
    section.top_margin = Inches(0)
    section.bottom_margin = Inches(0)
    section.left_margin = Inches(0)
    section.right_margin = Inches(0)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)

    paragraph = document.paragraphs[-1] if document.paragraphs else document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.add_run().add_picture(image_path, width=Inches(width_inches), height=Inches(height_inches))
    if not is_last:
        document.add_page_break()


def main():
    output_dir = os.path.dirname(OUTPUT)
    os.makedirs(output_dir, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="chaliyar_pages_", dir=output_dir) as image_dir:
        images = render_pages([FACING, PROJECT, GUIDE], image_dir)
        document = Document()
        document.core_properties.title = "Chaliyar River Pollution Project"
        document.core_properties.subject = "Dissertation project with facing sheets and guide biodata"
        document.core_properties.author = ""
        for index, (image_path, width, height) in enumerate(images):
            add_page(document, image_path, width, height, index == len(images) - 1)
        document.save(OUTPUT)
    print(OUTPUT)
    print(f"pages={len(images)}")


if __name__ == "__main__":
    main()
